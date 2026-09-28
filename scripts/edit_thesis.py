"""Apply the requested structural edits to the thesis .docx.

1. Move the quantitative sample-size / sampling-method subsections from
   Chapter 3 (3-3-2, 3-3-3) to the start of the quantitative analysis in
   Chapter 4 (new 4-2-1) and renumber the following 4-2-x headings.
2. Add a growth & learning breakdown by the three intangible capitals
   (human / information / organizational): a separate table and chart for
   each capital plus a comparison table and chart (new 4-2-3-6).

Usage: python3 edit_thesis.py <in.docx> <out.docx> <figures_dir>
"""
import copy
import os
import re
import sys

import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
src, dst, fig_dir = sys.argv[1:4]
doc = docx.Document(src)
body = doc.element.body
E = list(body.iterchildren())


def text_of(el):
    return "".join("-" if t.tag == W + "noBreakHyphen" else (t.text or "")
                   for t in el.iter(W + "t", W + "noBreakHyphen"))


def find(pred, start=0):
    for i in range(start, len(E)):
        if pred(E[i]):
            return i
    raise LookupError


def is_caption(e):
    ps = e.find(W + "pPr/" + W + "pStyle")
    return e.tag == W + "p" and ps is not None and ps.get(W + "val") == "Caption"


def find_caption(prefix):
    return find(lambda e: is_caption(e) and text_of(e).startswith(prefix))


def find_text(prefix, start=0):
    return find(lambda e: e.tag == W + "p" and text_of(e).strip().startswith(prefix), start)


def strip_bookmarks(el):
    for tag in ("bookmarkStart", "bookmarkEnd"):
        for b in list(el.iter(W + tag)):
            b.getparent().remove(b)
    return el


def make_run(template_run, text, bold=False):
    r = copy.deepcopy(template_run)
    for c in list(r):
        if c.tag != W + "rPr":
            r.remove(c)
    rpr = r.find(W + "rPr")
    if rpr is None:
        rpr = r.makeelement(W + "rPr", {})
        r.insert(0, rpr)
    if rpr.find(W + "rtl") is None:
        rpr.append(rpr.makeelement(W + "rtl", {}))
    if bold:
        for tag in ("b", "bCs"):
            rpr.insert(0, rpr.makeelement(W + tag, {}))
    t = r.makeelement(W + "t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
    t.text = text
    r.append(t)
    return r


def clone_para(template, texts, bold=False):
    """Copy a paragraph's formatting; replace its content by the given run texts."""
    p = strip_bookmarks(copy.deepcopy(template))
    runs = p.findall(W + "r")
    tmpl = runs[-1] if runs else p.makeelement(W + "r", {})
    for c in list(p):
        if c.tag != W + "pPr":
            p.remove(c)
    for t in texts:
        p.append(make_run(tmpl, t, bold))
    return p


def set_caption(p, number, title=None):
    """Set a caption's visible number and, optionally, its title text.

    The thesis captions show "<chapter>-<typed tens digit><SEQ last digit>", so the
    typed prefix before the SEQ field and the SEQ cached result are both updated.
    """
    runs = list(p.iter(W + "r"))
    kinds = [(r.find(W + "fldChar").get(W + "fldCharType") if r.find(W + "fldChar") is not None else None)
             for r in runs]
    seq_begin = max(i for i, k in enumerate(kinds) if k == "begin")
    seq_sep = next(i for i in range(seq_begin, len(runs)) if kinds[i] == "separate")
    seq_end = next(i for i in range(seq_sep, len(runs)) if kinds[i] == "end")
    hyphen = max(i for i in range(seq_begin) if runs[i].find(W + "noBreakHyphen") is not None)
    digits = str(number)
    prefix_ts = [t for r in runs[hyphen + 1:seq_begin] for t in r.iter(W + "t")]
    if prefix_ts:
        prefix_ts[0].text = digits[:-1]
        for t in prefix_ts[1:]:
            t.text = ""
    elif len(digits) > 1:
        runs[seq_begin].addprevious(make_run(runs[seq_sep + 1], digits[:-1]))
    seq_ts = [t for r in runs[seq_sep + 1:seq_end] for t in r.iter(W + "t")]
    seq_ts[0].text = digits[-1]
    for t in seq_ts[1:]:
        t.text = ""
    if title is not None:
        after = runs[seq_end + 1:]
        tmpl = after[0] if after else runs[seq_end]
        for r in after:
            r.getparent().remove(r)
        p.append(make_run(tmpl, title))


def replace_in_runs(p, old, new):
    for t in p.iter(W + "t"):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            return True
    raise LookupError(old)


def insert_after(anchor, els):
    for el in els:
        anchor.addnext(el)
        anchor = el
    return anchor


# ---------------------------------------------------------------- locate
ch3_sample_h = find_text("3-3-2-")               # حجم نمونه بخش کمّی
ch3_method_h = find_text("۳-۳-۳-")               # روش نمونه‌گیری
ch3_next_h = find_text("3-4-", ch3_method_h)     # تعاریف نظری متغیرها
qual_sample_p = ch3_sample_h - 1                 # body of 3-3-1
pop_p = find_text("توزیع این جامعه آماری")
q_intro = find_text("4-2- تجزیه‌وتحلیل کمّی")
h3_desc = find_text("4-2-1- تجزیه و تحلیل توصیفی")
h3_inf = find_text("4-2-2- تجزیه و تحلیل استنباطی")
h4_demo = find_text("4-2-1-1-")
body_p = h4_demo + 1
mixed_h = find_text("3-4-  تجزیه‌وتحلیل آمیخته")
rank_end = mixed_h - 1                           # closing paragraph of the ranking section
cap_tbl_tmpl = find_caption("جدول \u200f4-13")
cap_fig_tmpl = find_caption("شکل \u200f0-2میانگین وضعیت")
fig_para_tmpl = cap_fig_tmpl - 1
table_tmpl = find(lambda e: e.tag == W + "tbl" and "کشیدگی" in text_of(e))  # descriptive statistics
label_tmpl = find_text("حوزه اول: ذینفعان")
joint_cap = find_caption("جدول \u200f4-18تطبیقی")
integ_p = find(lambda e: "(جدول ۹-۴)" in text_of(e))

moved = E[ch3_sample_h:ch3_next_h]               # heading, text, caption, table, blank, heading, text, blank
old_sample_cap = [e for e in moved if text_of(e).startswith("جدول 1-3")][0]
sample_tbl = [e for e in moved if e.tag == W + "tbl"][0]
sample_text = moved[1]
method_text = [e for e in moved if text_of(e).startswith("روش نمونه‌گیری در این پژوهش")][0]

# --------------------------------------------- 1. remove from Chapter 3
for el in moved:
    body.remove(el)
insert_after(E[qual_sample_p], [clone_para(E[qual_sample_p], [
    "حجم نمونه و روش نمونه‌گیری بخش کمّی پژوهش، به‌منظور پیوستگی با تجزیه‌وتحلیل "
    "داده‌های پرسش‌نامه، در فصل چهارم (بند ۴-۲-۱) تشریح شده است."])])
replace_in_runs(E[pop_p], "که در بخش بعدی تشریح می‌شود", "که در فصل چهارم (بند ۴-۲-۱) تشریح شده است")

# ------------------------------- renumber existing 4-2-x headings (+1)
for i in range(h3_desc, mixed_h):
    el = E[i]
    ps = el.find(W + "pPr/" + W + "pStyle")
    if ps is None or ps.get(W + "val") not in ("Heading3", "Heading4"):
        continue
    t = next(el.iter(W + "t"))
    m = re.match(r"4-2-([12])", t.text)
    if m:
        t.text = t.text.replace(m.group(0), f"4-2-{int(m.group(1)) + 1}", 1)
t = next(E[mixed_h].iter(W + "t"))               # typo in the original: 3-4 -> 4-3
t.text = t.text.replace("3-4-", "4-3-", 1)

# --------------------------------------- 1b. insert sampling in Chapter 4
sample_cap = copy.deepcopy(E[cap_tbl_tmpl])
strip_bookmarks(sample_cap)
set_caption(sample_cap, 11, "حجم نمونه آماری بخش کمّی به تفکیک طبقات سازمانی")
blank = copy.deepcopy(moved[4])
insert_after(E[q_intro], [
    clone_para(E[h3_desc], ["4-2-1- ", "حجم نمونه و روش نمونه‌گیری بخش کمّی"]),
    clone_para(E[body_p], [
        "پیش از ارائه نتایج تحلیل داده‌های پرسش‌نامه، در این بخش حجم نمونه و روش "
        "نمونه‌گیری بخش کمّی پژوهش تشریح می‌شود؛ زیرا صد و هشت پاسخ تحلیل‌شده در ادامه "
        "این فصل، بر پایه همین طرح نمونه‌گیری گردآوری شده است."]),
    clone_para(E[h4_demo], ["4-2-1-1- ", "حجم نمونه بخش کمّی"]),
    sample_text, sample_cap, sample_tbl, blank,
    clone_para(E[h4_demo], ["4-2-1-2- ", "روش نمونه‌گیری بخش کمّی"]),
    method_text, copy.deepcopy(blank),
])

# shift cached numbers of the following Chapter 4 table captions (4-11..4-17 -> 4-12..4-18)
for el in E[h3_desc:mixed_h]:
    if is_caption(el) and "SEQ جدول" in "".join(x.text or "" for x in el.iter(W + "instrText")):
        m = re.match(r"جدول ‏4-(\d+)", text_of(el))
        set_caption(el, int(m.group(1)) + 1)

# ------------------------------------------ 2. growth & learning capitals
def fa_num(x, signed=False):
    s = f"{x:+.3f}" if signed else f"{x:.3f}"
    return s.translate(str.maketrans("0123456789.", "۰۱۲۳۴۵۶۷۸۹٫"))


def fa_int(n):
    return str(n).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def build_table(rows, widths):
    tbl = copy.deepcopy(E[table_tmpl])
    grid = tbl.find(W + "tblGrid")
    for g in list(grid):
        grid.remove(g)
    for w in widths:
        grid.append(grid.makeelement(W + "gridCol", {W + "w": str(w)}))
    trs = tbl.findall(W + "tr")
    head_tr, body_tr = trs[0], trs[1]
    for tr in trs:
        tbl.remove(tr)
    for ri, row in enumerate(rows):
        tr = copy.deepcopy(head_tr if ri == 0 else body_tr)
        tcs = tr.findall(W + "tc")
        tc_t = tcs[1] if len(tcs) > 1 else tcs[0]
        for tc in tcs:
            tr.remove(tc)
        for val in row:
            tc = copy.deepcopy(tc_t)
            p = tc.find(W + "p")
            runs = p.findall(W + "r")
            for r in runs:
                p.remove(r)
            p.append(make_run(runs[0], val, bold=(ri == 0 or row[0] == "—")))
            tr.append(tc)
        tbl.append(tr)
    return tbl


fig_no = [4]  # figures already in this section: scree, reliability, correlation, ranking


def figure(fname, caption):
    p = copy.deepcopy(E[fig_para_tmpl])
    for c in list(p):
        if c.tag != W + "pPr":
            p.remove(c)
    Paragraph(p, doc._body).add_run().add_picture(os.path.join(fig_dir, fname), width=Inches(5.9))
    cap = strip_bookmarks(copy.deepcopy(E[cap_fig_tmpl]))
    fig_no[0] += 1
    set_caption(cap, fig_no[0], caption)
    return [p, cap]


def table_caption(number, title):
    cap = strip_bookmarks(copy.deepcopy(E[cap_tbl_tmpl]))
    set_caption(cap, number, title)
    return cap


P = lambda text: clone_para(E[body_p], [text])
LABEL = lambda text: clone_para(E[label_tmpl], [text], bold=True)
MID, GL_MEAN = 3.0, 3.062
HEAD = ["ردیف", "مؤلفه", "گویه‌های مرتبط", "میانگین", "فاصله از حد متوسط (۳)", "رتبه در میان ۲۰ مؤلفه"]
WIDTHS = [700, 3560, 1300, 1100, 1500, 1200]

CAPITALS = [
    ("human", "سرمایه انسانی", [
        ("توسعه شایستگی‌های چندگانه تخصصی", "۲۱، ۲۲", 2.657, 1),
        ("آموزش تخصصی مبتنی بر شبیه‌ساز", "۲۵، ۲۶", 3.083, 5),
        ("مدیریت استعداد و جانشین‌پروری", "۲۷، ۲۸", 3.213, 10),
    ], "الف) سرمایه انسانی", [
        "جدول و شکل زیر وضعیت سه مؤلفه سرمایه انسانی را نشان می‌دهند. میانگین این سرمایه "
        "۹۸۴/۲ است که اندکی پایین‌تر از حد متوسط مقیاس (۳) و پایین‌تر از میانگین کل حوزه رشد "
        "و یادگیری (۰۶۲/۳) قرار دارد؛ بدین ترتیب، سرمایه انسانی تنها سرمایه‌ای است که میانگین "
        "آن از حد متوسط پایین‌تر است. این وضعیت عمدتاً ناشی از مؤلفه «توسعه شایستگی‌های چندگانه "
        "تخصصی» با میانگین ۶۵۷/۲ است که پایین‌ترین میانگین را در میان کل بیست مؤلفه پژوهش دارد "
        "(رتبه نخست در اولویت بهبود). در مقابل، دو مؤلفه «آموزش تخصصی مبتنی بر شبیه‌ساز» "
        "(۰۸۳/۳) و «مدیریت استعداد و جانشین‌پروری» (۲۱۳/۳) اندکی بالاتر از حد متوسط قرار دارند.",
        "اختلاف ۵۵۶/۰ واحدی میان بالاترین و پایین‌ترین مؤلفه این سرمایه نشان می‌دهد که ضعف "
        "سرمایه انسانی، بیش از آنکه به سازوکارهای آموزش شبیه‌سازی‌شده و مدیریت استعداد مربوط "
        "باشد، از شکاف شایستگی‌های فنی، شناختی و دیجیتالی کارکنان در کار با سامانه‌های نوین و "
        "ضعف آموزش تفکر سیستمی (گویه‌های ۲۱ و ۲۲) سرچشمه می‌گیرد.",
    ]),
    ("information", "سرمایه اطلاعاتی", [
        ("نظام یادگیری سازمانی پویا و مدیریت دانش دفاعی", "۲۳، ۲۴", 3.208, 9),
    ], "ب) سرمایه اطلاعاتی", [
        "سرمایه اطلاعاتی در این پژوهش از طریق مؤلفه «نظام یادگیری سازمانی پویا و مدیریت دانش "
        "دفاعی» سنجیده شده است که ثبت و انتقال نظام‌مند تجربیات عملیاتی و درس‌آموخته‌های جنگ‌های "
        "پیشین (گویه ۲۳) و سرعت انطباق ساختار آموزشی با تغییرات تهدیدات و فناوری‌های نوین "
        "(گویه ۲۴) را در بر می‌گیرد. میانگین این سرمایه ۲۰۸/۳ است که بالاترین میانگین را در میان "
        "سه سرمایه حوزه رشد و یادگیری دارد و ۱۴۶/۰ واحد از میانگین کل حوزه بالاتر است. با این حال، "
        "این مقدار فاصله اندکی با حد متوسط مقیاس دارد و مؤلفه مذکور در رتبه نهم اولویت بهبود در "
        "میان بیست مؤلفه قرار گرفته است؛ از این رو، وضعیت سرمایه اطلاعاتی را می‌توان «متوسط رو به "
        "بالا» و نه مطلوب ارزیابی کرد.",
    ]),
    ("organizational", "سرمایه سازمانی", [
        ("نوآوری، دکترین نبرد ناهمتراز و فرهنگ یادگیری", "۲۹، ۳۰", 3.148, 6),
    ], "ج) سرمایه سازمانی", [
        "سرمایه سازمانی از طریق مؤلفه «نوآوری، دکترین نبرد ناهمتراز و فرهنگ یادگیری» سنجیده "
        "شده است که تشویق و حمایت از خلاقیت و نوآوری پرسنل در ارائه راه‌حل‌های پدافندی (گویه ۲۹) "
        "و گسترش آموزش کاربرد هوش مصنوعی و فناوری‌های نوین (گویه ۳۰) را در بر می‌گیرد. میانگین "
        "این سرمایه ۱۴۸/۳ است که ۱۴۸/۰ واحد بالاتر از حد متوسط مقیاس و ۰۸۶/۰ واحد بالاتر از "
        "میانگین کل حوزه قرار دارد؛ اما این مؤلفه در رتبه ششم اولویت بهبود در میان بیست مؤلفه "
        "پژوهش جای گرفته است. به بیان دیگر، فرهنگ نوآوری و یادگیری، هرچند در وضعیتی بهتر از "
        "سرمایه انسانی قرار دارد، همچنان یکی از حلقه‌های نیازمند تقویت در الگوی توسعه متوازن "
        "دفاع هوایی به شمار می‌رود.",
    ]),
]

new = [
    clone_para(E[h4_demo], ["4-2-3-6- ", "تحلیل تفکیکی حوزه رشد و یادگیری بر پایه سه سرمایه نامشهود (انسانی، اطلاعاتی و سازمانی)"]),
    P("همان‌گونه که در بند ۲-۵-۴-۳ فصل دوم تشریح شد، کاپلان و نورتون (Kaplan & Norton, 2004) "
      "منظر رشد و یادگیری را در قالب سه دسته از دارایی‌های نامشهود، یعنی سرمایه انسانی، سرمایه "
      "اطلاعاتی و سرمایه سازمانی، صورت‌بندی کرده‌اند. از آنجا که حوزه رشد و یادگیری با میانگین "
      "۰۶۲/۳ پایین‌ترین وضعیت موجود را در میان حوزه‌های چهارگانه داشت و تفاوت آن با حد متوسط "
      "مقیاس نیز معنادار نبود، برای شناسایی دقیق‌تر خاستگاه این ضعف، پنج مؤلفه این حوزه بر پایه "
      "محتوای مفهومی خود به سه سرمایه یادشده نسبت داده شد و وضعیت هر سرمایه به‌طور جداگانه در "
      "قالب جدول و نمودار مستقل بررسی شد."),
    P("بر این اساس، سه مؤلفه «توسعه شایستگی‌های چندگانه تخصصی» (گویه‌های ۲۱ و ۲۲)، «آموزش "
      "تخصصی مبتنی بر شبیه‌ساز» (گویه‌های ۲۵ و ۲۶) و «مدیریت استعداد و جانشین‌پروری» (گویه‌های "
      "۲۷ و ۲۸) که به دانش، مهارت و شایستگی‌های کارکنان مربوط‌اند، ذیل سرمایه انسانی؛ مؤلفه "
      "«نظام یادگیری سازمانی پویا و مدیریت دانش دفاعی» (گویه‌های ۲۳ و ۲۴) که به ثبت، نگهداشت و "
      "انتقال دانش و تجربیات عملیاتی مربوط است، ذیل سرمایه اطلاعاتی؛ و مؤلفه «نوآوری، دکترین "
      "نبرد ناهمتراز و فرهنگ یادگیری» (گویه‌های ۲۹ و ۳۰) که به فرهنگ، ارزش‌ها و ظرفیت تغییر "
      "سازمان مربوط است، ذیل سرمایه سازمانی قرار گرفت. میانگین مؤلفه‌ها از جدول رتبه‌بندی بیست "
      "مؤلفه (جدول ۱۸-۴) برگرفته شده و میانگین هر سرمایه، میانگین ساده مؤلفه‌های آن است؛ از آنجا "
      "که هر مؤلفه با دو گویه سنجیده شده، این مقدار با میانگین گویه‌های همان سرمایه برابر است."),
]

tbl_no = 18
summary = []
for key, title, comps, label, paras in CAPITALS:
    mean = round(sum(c[2] for c in comps) / len(comps), 3)
    summary.append((title, len(comps), mean))
    rows = [HEAD] + [[fa_int(i + 1), n, items, fa_num(v), fa_num(v - MID, True), fa_int(rk)]
                     for i, (n, items, v, rk) in enumerate(comps)]
    rows.append(["—", f"میانگین {title}", f"{fa_int(2 * len(comps))} گویه", fa_num(mean),
                 fa_num(mean - MID, True), "—"])
    tbl_no += 1
    new += [LABEL(label)] + [P(t) for t in paras] + [
        table_caption(tbl_no, f"وضعیت موجود مؤلفه‌های {title} در حوزه رشد و یادگیری (مقیاس ۱ تا ۵، N=108)"),
        build_table(rows, WIDTHS), copy.deepcopy(blank),
    ] + figure(f"gl_{key}_capital.png", f"میانگین مؤلفه‌های {title} در حوزه رشد و یادگیری")

order = sorted(summary, key=lambda s: s[2])
prio = {t: i + 1 for i, (t, _, _) in enumerate(order)}
rows = [["سرمایه", "تعداد مؤلفه", "تعداد گویه", "میانگین", "فاصله از حد متوسط (۳)", "اولویت بهبود"]]
rows += [[t, fa_int(n), fa_int(2 * n), fa_num(m), fa_num(m - MID, True), fa_int(prio[t])] for t, n, m in summary]
rows.append(["—", fa_int(5), fa_int(10), fa_num(GL_MEAN), fa_num(GL_MEAN - MID, True), "—"])
rows[-1][0] = "کل حوزه رشد و یادگیری"
tbl_no += 1
new += [
    LABEL("د) مقایسه سه سرمایه نامشهود"),
    P("مقایسه سه سرمایه (جدول و شکل زیر) نشان می‌دهد که ترتیب اولویت بهبود در حوزه رشد و "
      "یادگیری به‌صورت سرمایه انسانی (۹۸۴/۲)، سرمایه سازمانی (۱۴۸/۳) و سرمایه اطلاعاتی (۲۰۸/۳) "
      "است. به بیان دیگر، پایین بودن میانگین کل حوزه رشد و یادگیری و معنادار نبودن تفاوت آن با "
      "حد متوسط، بیش از همه از ضعف سرمایه انسانی و به‌ویژه شکاف شایستگی‌های تخصصی کارکنان "
      "سرچشمه می‌گیرد. این یافته با دیدگاه کاپلان و نورتون مبنی بر زیربنایی بودن سرمایه انسانی "
      "برای تحقق اهداف سایر منظرها همسوست و نشان می‌دهد که مداخلات مدیریتی در این حوزه باید "
      "پیش از هر چیز بر توسعه شایستگی‌های تخصصی کارکنان متمرکز شود."),
    table_caption(tbl_no, "مقایسه وضعیت موجود سه سرمایه نامشهود حوزه رشد و یادگیری"),
    build_table(rows, [2400, 1150, 1150, 1300, 1900, 1460]),
    copy.deepcopy(blank),
] + figure("gl_capitals_comparison.png", "مقایسه میانگین سه سرمایه نامشهود حوزه رشد و یادگیری") + [
    P("لازم به ذکر است که این مقایسه ماهیتی توصیفی دارد؛ زیرا سرمایه‌های اطلاعاتی و سازمانی هر "
      "یک تنها با یک مؤلفه (دو گویه) سنجیده شده‌اند و تفاوت میانگین‌ها نیز کوچک است. بنابراین، "
      "داوری درباره معناداری آماری تفاوت میان سه سرمایه، نیازمند آزمون‌های تکمیلی بر پایه "
      "داده‌های خام پاسخ‌دهندگان است."),
]
insert_after(E[rank_end], new)

# joint-display table (typed number, no SEQ field) and in-text references
tbl_no += 1
replace_in_runs(E[joint_cap], "18", str(tbl_no))
replace_in_runs(E[integ_p], "(جدول ۹-۴)", "(جدول ۱۸-۴)")
replace_in_runs(E[integ_p], "(جدول ۸-۴)", "(جدول ۱۷-۴)")

# copied paragraphs/rows carry the template's w14:paraId/textId; keep them unique
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
seen = set()
for el in body.iter(W + "p", W + "tr"):
    pid = el.get(W14 + "paraId")
    if pid is None:
        continue
    if pid in seen:
        del el.attrib[W14 + "paraId"]
        el.attrib.pop(W14 + "textId", None)
    else:
        seen.add(pid)

# new pictures were created outside the tree, so they share one docPr id
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
docprs = list(body.iter(WP + "docPr"))
next_id = max(int(d.get("id")) for d in docprs) + 1
used = set()
for d in docprs:
    if d.get("id") in used:
        d.set("id", str(next_id))
        d.set("name", f"Picture {next_id}")
        next_id += 1
    used.add(d.get("id"))

doc.save(dst)
print("saved", dst, "last table number", tbl_no)
